"""Four deliberately awkward Word documents for testing DocForge's importer.

Generated with python-docx rather than the project's own corpus generator, so
the importer meets OOXML it has not been tuned against.
"""
import io
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).parent


def png(width, height, rgb):
	"""A solid-colour PNG, built by hand so nothing needs Pillow."""
	row = b"\x00" + bytes(rgb) * width
	raw = row * height

	def chunk(kind, data):
		body = kind + data
		return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

	return io.BytesIO(
		b"\x89PNG\r\n\x1a\n"
		+ chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
		+ chunk(b"IDAT", zlib.compress(raw, 9))
		+ chunk(b"IEND", b"")
	)


def shade(cell, hex_fill):
	tc_pr = cell._tc.get_or_add_tcPr()
	shd = OxmlElement("w:shd")
	shd.set(qn("w:val"), "clear")
	shd.set(qn("w:color"), "auto")
	shd.set(qn("w:fill"), hex_fill)
	tc_pr.append(shd)


def hyperlink(paragraph, text, url):
	rel = paragraph.part.relate_to(
		url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
	)
	link = OxmlElement("w:hyperlink")
	link.set(qn("r:id"), rel)
	run = OxmlElement("w:r")
	r_pr = OxmlElement("w:rPr")
	colour = OxmlElement("w:color")
	colour.set(qn("w:val"), "0563C1")
	under = OxmlElement("w:u")
	under.set(qn("w:val"), "single")
	r_pr.append(colour)
	r_pr.append(under)
	run.append(r_pr)
	t = OxmlElement("w:t")
	t.text = text
	run.append(t)
	link.append(run)
	paragraph._p.append(link)


def field(paragraph, instruction, placeholder):
	"""A complex field (TOC, PAGE, DATE): begin, instruction, separate, result, end."""
	for kind in ("begin", None, "separate", "result", "end"):
		run = paragraph.add_run()
		if kind == "result":
			run.text = placeholder
			continue
		if kind is None:
			instr = OxmlElement("w:instrText")
			instr.set(qn("xml:space"), "preserve")
			instr.text = instruction
			run._r.append(instr)
			continue
		char = OxmlElement("w:fldChar")
		char.set(qn("w:fldCharType"), kind)
		run._r.append(char)


def tracked_insert(paragraph, text, author="Reviewer"):
	ins = OxmlElement("w:ins")
	ins.set(qn("w:id"), "901")
	ins.set(qn("w:author"), author)
	ins.set(qn("w:date"), "2026-09-01T09:00:00Z")
	run = OxmlElement("w:r")
	t = OxmlElement("w:t")
	t.set(qn("xml:space"), "preserve")
	t.text = text
	run.append(t)
	ins.append(run)
	paragraph._p.append(ins)


def tracked_delete(paragraph, text, author="Reviewer"):
	dele = OxmlElement("w:del")
	dele.set(qn("w:id"), "902")
	dele.set(qn("w:author"), author)
	dele.set(qn("w:date"), "2026-09-01T09:05:00Z")
	run = OxmlElement("w:r")
	t = OxmlElement("w:delText")
	t.set(qn("xml:space"), "preserve")
	t.text = text
	run.append(t)
	dele.append(run)
	paragraph._p.append(dele)


# ---------------------------------------------------------------- document 1
def board_risk_report():
	doc = Document()
	section = doc.sections[0]
	section.header.paragraphs[0].text = "Group Risk Committee  |  Quarterly Risk Report  |  CONFIDENTIAL"
	section.footer.paragraphs[0].text = "Prepared by Enterprise Risk Management. Not for onward distribution."

	title = doc.add_heading("Quarterly Enterprise Risk Report", 0)
	title.alignment = WD_ALIGN_PARAGRAPH.CENTER
	sub = doc.add_paragraph()
	sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
	run = sub.add_run("Third quarter, financial year 2026")
	run.italic = True
	run.font.size = Pt(13)
	run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

	doc.add_heading("1. Executive summary", 1)
	p = doc.add_paragraph()
	p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
	p.add_run("The overall risk profile is ")
	r = p.add_run("stable with two items outside appetite")
	r.bold = True
	p.add_run(". Operational resilience improved following the recovery exercise, while ")
	r = p.add_run("third-party concentration")
	r.italic = True
	r.underline = True
	p.add_run(" remains the principal concern. The previous rating of ")
	r = p.add_run("amber")
	r.font.strike = True
	p.add_run(" has been revised to ")
	r = p.add_run(" RED ")
	r.bold = True
	r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
	r.font.highlight_color = WD_COLOR_INDEX.RED
	p.add_run(". Capital headroom is 1.4x the regulatory minimum (CET1")
	r = p.add_run("a")
	r.font.superscript = True
	p.add_run("), and liquidity H")
	r = p.add_run("2")
	r.font.subscript = True
	p.add_run(" projections are unchanged.")

	doc.add_heading("2. Risk heat map", 1)
	table = doc.add_table(rows=6, cols=6)
	table.style = "Table Grid"
	table.alignment = WD_TABLE_ALIGNMENT.CENTER
	heads = ["Risk category", "Owner", "Inherent", "Residual", "Trend", "Appetite"]
	for i, text in enumerate(heads):
		cell = table.rows[0].cells[i]
		cell.text = ""
		run = cell.paragraphs[0].add_run(text)
		run.bold = True
		run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
		shade(cell, "1F3864")
	rows = [
		("Credit", "Chief Credit Officer", "High", "Medium", "Stable", "Within"),
		("Operational", "Chief Operating Officer", "High", "Medium", "Improving", "Within"),
		("Third party", "Chief Procurement Officer", "Critical", "High", "Worsening", "OUTSIDE"),
		("Cyber", "Chief Information Security Officer", "Critical", "High", "Stable", "OUTSIDE"),
		("Conduct", "Chief Compliance Officer", "Medium", "Low", "Stable", "Within"),
	]
	fills = {"Critical": "C00000", "High": "ED7D31", "Medium": "FFD966", "Low": "A9D18E", "OUTSIDE": "C00000", "Within": "A9D18E"}
	for r_i, row in enumerate(rows, start=1):
		for c_i, text in enumerate(row):
			cell = table.rows[r_i].cells[c_i]
			cell.text = text
			if text in fills:
				shade(cell, fills[text])
				cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

	doc.add_heading("3. Items outside appetite", 1)
	doc.add_heading("3.1 Third-party concentration", 2)
	for text in (
		"Four critical services depend on a single hosting provider.",
		"Exit plans exist for two of the four; the remainder are due next quarter.",
		"Contractual audit rights have not been exercised in 24 months.",
	):
		doc.add_paragraph(text, style="List Bullet")
	doc.add_heading("3.2 Agreed actions", 2)
	for text in (
		"Complete exit plans for the remaining two services.",
		"Exercise audit rights and report findings to this committee.",
		"Commission an independent concentration study.",
	):
		doc.add_paragraph(text, style="List Number")

	doc.add_heading("4. Exposure by region", 1)
	merged = doc.add_table(rows=5, cols=4)
	merged.style = "Table Grid"
	top = merged.cell(0, 0).merge(merged.cell(0, 3))
	top.text = "Gross exposure by region and segment (millions)"
	top.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
	top.paragraphs[0].runs[0].bold = True
	shade(top, "D9E2F3")
	for i, text in enumerate(["Region", "Retail", "Commercial", "Total"]):
		merged.cell(1, i).text = text
		merged.cell(1, i).paragraphs[0].runs[0].bold = True
	side = merged.cell(2, 0).merge(merged.cell(3, 0))
	side.text = "Domestic (two books)"
	for r_i, values in ((2, ("412.5", "1,208.0", "1,620.5")), (3, ("88.1", "240.9", "329.0"))):
		for c_i, text in enumerate(values, start=1):
			merged.cell(r_i, c_i).text = text
			merged.cell(r_i, c_i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
	merged.cell(4, 0).text = "International"
	for c_i, text in enumerate(("57.3", "612.4", "669.7"), start=1):
		merged.cell(4, c_i).text = text
		merged.cell(4, c_i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

	doc.add_heading("5. Trend chart", 1)
	doc.add_picture(png(600, 160, (31, 56, 100)), width=Inches(6))
	caption = doc.add_paragraph("Figure 1. Placeholder for the residual risk trend.")
	caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
	caption.runs[0].italic = True

	doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
	doc.add_heading("Appendix A. Committee statement", 1)
	doc.add_paragraph(
		"The committee notes the position and requires a remediation plan for both items "
		"outside appetite to be presented at its next meeting.",
		style="Intense Quote",
	)
	doc.save(OUT / "01-board-risk-report.docx")


# ---------------------------------------------------------------- document 2
def policy_with_nested_tables():
	doc = Document()
	doc.sections[0].header.paragraphs[0].text = "Information Security Policy  |  Version 4.2"
	doc.sections[0].footer.paragraphs[0].text = "Controlled document. Printed copies are uncontrolled."
	doc.add_heading("Information Security Policy", 0)

	doc.add_heading("Document control", 1)
	control = doc.add_table(rows=4, cols=2)
	control.style = "Table Grid"
	for i, (k, v) in enumerate(
		[("Owner", "Chief Information Security Officer"), ("Approved by", "Group Risk Committee"),
		 ("Review cycle", "Annual"), ("Classification", "Internal")]
	):
		control.cell(i, 0).text = k
		control.cell(i, 0).paragraphs[0].runs[0].bold = True
		shade(control.cell(i, 0), "F2F2F2")
		control.cell(i, 1).text = v

	doc.add_heading("1. Purpose and scope", 1)
	p = doc.add_paragraph("This policy applies to all staff, contractors and third parties. Supporting standards are published at ")
	hyperlink(p, "the internal policy library", "https://policies.example.invalid/library")
	p.add_run(" and must be read alongside it.")

	doc.add_heading("2. Control requirements", 1)
	doc.add_heading("2.1 Access control", 2)
	doc.add_heading("2.1.1 Privileged access", 3)
	doc.add_heading("Break-glass accounts", 4)
	doc.add_paragraph("Break-glass accounts are sealed, monitored and reviewed after every use.")

	doc.add_heading("2.2 Multi-level obligations", 2)
	levels = [
		("List Number", "Identity"),
		("List Number 2", "Joiners are provisioned from the system of record only."),
		("List Number 2", "Leavers are disabled within one working day."),
		("List Number 3", "Privileged leavers within four hours."),
		("List Number", "Cryptography"),
		("List Bullet 2", "Approved algorithms only."),
		("List Bullet 3", "Key lengths reviewed annually."),
		("List Number", "Logging"),
	]
	for style, text in levels:
		doc.add_paragraph(text, style=style)

	doc.add_heading("3. Responsibilities (nested table)", 1)
	outer = doc.add_table(rows=2, cols=2)
	outer.style = "Table Grid"
	outer.cell(0, 0).text = "First line"
	outer.cell(0, 1).text = "Second line"
	for c in (0, 1):
		outer.cell(0, c).paragraphs[0].runs[0].bold = True
		shade(outer.cell(0, c), "E2EFDA")
	outer.cell(1, 0).text = "Operate the controls and evidence them."
	host = outer.cell(1, 1)
	host.text = "Oversee and challenge. Testing schedule:"
	inner = host.add_table(rows=3, cols=2)
	inner.style = "Table Grid"
	for i, (a, b) in enumerate([("Control", "Frequency"), ("Access recertification", "Quarterly"), ("Log review", "Monthly")]):
		inner.cell(i, 0).text = a
		inner.cell(i, 1).text = b
	shade(inner.cell(0, 0), "FFF2CC")
	shade(inner.cell(0, 1), "FFF2CC")

	doc.add_heading("4. Typography stress", 1)
	p = doc.add_paragraph()
	for font, size in (("Georgia", 14), ("Courier New", 10), ("Arial", 12), ("Times New Roman", 16)):
		r = p.add_run(f"{font} at {size} point.  ")
		r.font.name = font
		r.font.size = Pt(size)
	p = doc.add_paragraph()
	for colour, word in ((RGBColor(0xC0, 0, 0), "prohibited "), (RGBColor(0xBF, 0x8F, 0), "restricted "), (RGBColor(0x37, 0x86, 0x3D), "permitted")):
		r = p.add_run(word)
		r.bold = True
		r.font.color.rgb = colour
	p = doc.add_paragraph()
	r = p.add_run("Highlighted obligation: exceptions require written approval.")
	r.font.highlight_color = WD_COLOR_INDEX.YELLOW
	right = doc.add_paragraph("Right-aligned sign-off block")
	right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
	doc.save(OUT / "02-policy-nested-tables.docx")


# ---------------------------------------------------------------- document 3
def unsupported_features():
	"""Everything the fidelity notes say is dropped: the question is whether it
	is dropped cleanly, with the surrounding text intact."""
	doc = Document()
	doc.add_heading("Contract Review with Tracked Changes", 0)

	doc.add_heading("Table of contents (field)", 1)
	field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "Right-click to update the table of contents.")

	doc.add_heading("1. Clause under negotiation", 1)
	p = doc.add_paragraph("The supplier shall deliver the services within ")
	tracked_delete(p, "thirty (30)")
	tracked_insert(p, "fourteen (14)")
	p.add_run(" days of the effective date. TEXT-AFTER-TRACKED-CHANGE must survive.")

	doc.add_heading("2. Fields inside running text", 1)
	p = doc.add_paragraph("This page is number ")
	field(p, "PAGE", "1")
	p.add_run(" and was printed on ")
	field(p, 'DATE \\@ "d MMMM yyyy"', "18 September 2026")
	p.add_run(". TEXT-AFTER-FIELDS must survive.")

	doc.add_heading("3. Tabs, breaks and non-breaking characters", 1)
	p = doc.add_paragraph("Name:\tValue one\tValue two")
	p = doc.add_paragraph("First line of an address")
	p.add_run().add_break(WD_BREAK.LINE)
	p.add_run("Second line after a soft return")
	doc.add_paragraph("Non‑breaking hyphen, non breaking space, soft­hyphen, en–dash, em—dash, ellipsis…")

	doc.add_heading("4. Landscape section follows", 1)
	doc.add_paragraph("The next section switches to landscape with a wide table.")
	wide = doc.add_section(WD_SECTION.NEW_PAGE)
	wide.orientation = WD_ORIENT.LANDSCAPE
	wide.page_width, wide.page_height = wide.page_height, wide.page_width
	doc.add_heading("5. Wide schedule (second section, landscape)", 1)
	table = doc.add_table(rows=4, cols=10)
	table.style = "Table Grid"
	for c in range(10):
		table.cell(0, c).text = f"Month {c + 1}"
		shade(table.cell(0, c), "DDEBF7")
		for r in range(1, 4):
			table.cell(r, c).text = f"{(r * 137 + c * 53) % 1000:,}.00"
	doc.add_paragraph("TEXT-AFTER-SECOND-SECTION must survive.")
	doc.save(OUT / "03-tracked-changes-fields-sections.docx")


# ---------------------------------------------------------------- document 4
def multilingual_and_scale():
	doc = Document()
	doc.add_heading("Multilingual Notice and Large Register", 0)
	doc.add_heading("1. Notice in several scripts", 1)
	for label, text in (
		("English", "All employees must complete the annual attestation."),
		("Arabic", "يجب على جميع الموظفين إكمال الإقرار السنوي."),
		("Chinese", "所有员工必须完成年度声明。"),
		("Japanese", "全従業員は年次証明を完了する必要があります。"),
		("Hindi", "सभी कर्मचारियों को वार्षिक सत्यापन पूरा करना होगा।"),
		("Symbols", "€ 1,250.00  £ 980.50  ¥ 140,000  ≤ ≥ ± µ ™ ©  ✓ ✗  emoji: \U0001F512 \U0001F4C4"),
		("Markup-like", "<script>alert('x')</script> & \"quotes\" <b>not bold</b>"),
	):
		p = doc.add_paragraph()
		p.add_run(f"{label}: ").bold = True
		p.add_run(text)

	doc.add_heading("2. Risk register (120 rows, banded)", 1)
	rows = 120
	table = doc.add_table(rows=rows + 1, cols=5)
	table.style = "Table Grid"
	for c, text in enumerate(["Ref", "Risk", "Owner", "Rating", "Due"]):
		cell = table.cell(0, c)
		cell.text = text
		cell.paragraphs[0].runs[0].bold = True
		shade(cell, "203864")
		cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
	ratings = [("Low", "C6E0B4"), ("Medium", "FFE699"), ("High", "F4B084"), ("Critical", "FF7C80")]
	for r in range(1, rows + 1):
		rating, fill = ratings[(r * 7) % 4]
		values = (f"R-{r:04d}", f"Register entry {r}: control weakness in process area {r % 13 + 1}",
		          f"Owner {r % 9 + 1}", rating, f"2026-{(r % 12) + 1:02d}-{(r % 27) + 1:02d}")
		for c, text in enumerate(values):
			table.cell(r, c).text = text
		shade(table.cell(r, 3), fill)
	doc.add_paragraph("END-OF-REGISTER marker paragraph.")

	doc.add_heading("3. Several pictures", 1)
	for rgb in ((192, 0, 0), (0, 112, 192), (0, 176, 80)):
		doc.add_picture(png(240, 60, rgb), width=Inches(2.5))
	doc.save(OUT / "04-multilingual-large-register.docx")


for build in (board_risk_report, policy_with_nested_tables, unsupported_features, multilingual_and_scale):
	build()
for f in sorted(OUT.glob("*.docx")):
	print(f"{f.name}  {f.stat().st_size:,} bytes")
